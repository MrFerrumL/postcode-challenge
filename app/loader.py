"""Чтение блоков одного листа. Позиции ищем по заголовкам, связи — по ID."""
import math
import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import openpyxl
from .schemas import Dataset, Engineer, History, Task


def role_name(value):
    return " ".join(str(value).strip().casefold().split()).replace("big data", "bigdata")


def skill_name(value):
    return " ".join(str(value).strip().casefold().split())


def number(value, context):
    try:
        result = float(str(value).replace(",", "."))
    except (ValueError, TypeError):
        raise ValueError(f"{context}: требуется число, получено {value!r}")
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"{context}: число должно быть конечным и неотрицательным")
    return result


def day(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def load_excel(content: bytes) -> Dataset:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50_000_000:
                raise ValueError("Распакованный Excel превышает 50 МБ")
        book = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Не удалось прочитать XLSX") from exc
    if len(book.sheetnames) != 1:
        raise ValueError("Ожидается один лист с блоками исходного датасета")
    sheet = book.active
    if sheet.max_row > 20000 or sheet.max_column > 1000:
        raise ValueError("Слишком большой лист")
    rows = list(sheet.values)
    book.close()

    def block(required):
        for n, row in enumerate(rows):
            if set(required).issubset(set(row)):
                headers = [str(x).strip() if x is not None else None for x in row]
                records = []
                for following in rows[n+1:]:
                    if following[0] is None:
                        break
                    records.append({h: following[i] for i, h in enumerate(headers) if h})
                return records
        raise ValueError(f"Не найден блок: {', '.join(required)}")

    raw_tasks = block(["task_id", "team_id", "estimation_sp"])
    raw_spent = block(["task_id", "Исходное время"])
    raw_people = block(["engineer_id", "capacity_rate"])
    raw_history = block(["snapshot_date", "velocity_achieved"])
    raw_deps = block(["blocking_task_id", "blocked_task_id"])
    raw_estimates = block(["Роль"])
    estimates = defaultdict(dict)
    totals = {}
    for record in raw_estimates:
        label = role_name(record["Роль"])
        for tid, value in record.items():
            if tid == "Роль":
                continue
            hours = number(value, f"Смета {tid}, {label}")
            if label.startswith("итого"):
                totals[tid] = hours
            else:
                estimates[tid][label] = estimates[tid].get(label, 0) + hours

    spent = {}
    for record in raw_spent:
        tid = record["task_id"]
        if tid in spent:
            raise ValueError(f"Повтор spent_time_roles: {tid}")
        spent[tid] = {}
        for key, value in record.items():
            if key not in ("task_id", "Исходное время"):
                role = role_name(key)
                spent[tid][role] = spent[tid].get(role, 0) + number(value, f"Факт {tid}, {role}")

    issues = []
    def issue(code, message, task=None):
        issues.append({"code": code, "message": message, "task_id": task})

    tasks = []
    for r in raw_tasks:
        tid = str(r["task_id"])
        if tid not in estimates:
            raise ValueError(f"Нет сметы для {tid}")
        task = Task(id=tid, initiative=str(r["Номер инициативы"]), title=str(r["summary"]),
                    team=str(r["team_id"]), status=r["status"], priority=number(r["rung"], tid),
                    sp=number(r["estimation_sp"], tid), estimated_hours=number(r["estimated_hh"], tid),
                    roles=estimates[tid], spent=spent.get(tid, {}), created=day(r["created_at"]),
                    actual_end=day(r["actual_end"]))
        total = sum(task.roles.values())
        if abs(total-task.estimated_hours) > 1e-6:
            issue("hours_mismatch", f"{tid}: Tasks {task.estimated_hours:g} ч; сумма ролей {total:g} ч", tid)
        if tid in totals and abs(total-totals[tid]) > 1e-6:
            issue("total_mismatch", f"{tid}: строка итога {totals[tid]:g} ч; сумма ролей {total:g} ч", tid)
        for role, used in task.spent.items():
            if used > task.roles.get(role, 0):
                issue("overrun", f"{tid}: {role}, потрачено {used:g}, оценка {task.roles.get(role, 0):g} ч", tid)
        if day(r["planned_start"]) and task.created and day(r["planned_start"]) < task.created:
            issue("date_conflict", f"{tid}: плановый старт раньше создания", tid)
        tasks.append(task)

    engineers = [Engineer(id=str(r["engineer_id"]), team=str(r["team_id"]), role=role_name(r["role"]),
                         grade=str(r["grade"]), rate=number(r["capacity_rate"], str(r["engineer_id"])),
                         skills=[skill_name(x) for x in re.split(r",(?![^()]*\))", str(r["skills_declared"])) if x.strip()])
                 for r in raw_people]
    history = [History(date=day(r["snapshot_date"]), team=str(r["team_id"]),
                       achieved=number(r["velocity_achieved"], "velocity"), planned=number(r["planned_sp"], "planned_sp"))
               for r in raw_history]
    dependencies = [(str(r["blocking_task_id"]), str(r["blocked_task_id"])) for r in raw_deps]
    dataset = Dataset(tasks=tasks, engineers=engineers, history=history, dependencies=dependencies, issues=issues)
    validate_dataset(dataset)
    return dataset


def validate_dataset(data):
    from graphlib import TopologicalSorter, CycleError
    ids = {t.id for t in data.tasks}
    if not ids or len(ids) != len(data.tasks):
        raise ValueError("Пустой набор задач или повторяющиеся task_id")
    rates = defaultdict(float)
    assignments = set()
    profiles = {}
    for e in data.engineers:
        if (e.id, e.team) in assignments:
            raise ValueError(f"Дублируется назначение {e.id}, {e.team}")
        assignments.add((e.id, e.team))
        rates[e.id] += e.rate
        profile = (e.role, e.grade, tuple(sorted(e.skills)))
        if e.id in profiles and profiles[e.id] != profile:
            raise ValueError(f"Противоречащие профили {e.id}")
        profiles[e.id] = profile
    if any(v > 1 + 1e-8 for v in rates.values()):
        raise ValueError("Суммарная доля сотрудника больше 1")
    keys = [(h.team, h.date) for h in data.history]
    if len(keys) != len(set(keys)):
        raise ValueError("Дублируется команда/дата в истории")
    graph = {tid: set() for tid in ids}
    tasks = {t.id:t for t in data.tasks}
    for a,b in data.dependencies:
        if a not in ids or b not in ids:
            raise ValueError(f"Неизвестная задача зависимости {a} → {b}")
        graph[b].add(a)
        if tasks[b].status == "Done" and tasks[a].status != "Done":
            data.issues.append({"code":"dependency_fact_conflict", "task_id":b,
                                "message":f"{a} → {b}: предшественник не Done, последующая уже Done"})
    try:
        list(TopologicalSorter(graph).static_order())
    except CycleError as exc:
        raise ValueError("Цикл зависимостей: " + str(exc.args[1])) from exc
